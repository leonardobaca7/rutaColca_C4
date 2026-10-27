import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_report():
    doc = docx.Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        # Header / Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Construcción de Software (2026-B) · EPIS UNSA | Caso RutaColca (AG-08 / AG-09)")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(9)
        f_run.font.color.rgb = RGBColor(120, 120, 120)

    FONT_FAMILY = "Times New Roman"
    COLOR_PRIMARY = RGBColor(11, 61, 98)      # #0B3D62
    COLOR_SECONDARY = RGBColor(27, 94, 139)   # #1B5E8B
    COLOR_TEXT = RGBColor(20, 20, 20)
    COLOR_MUTED = RGBColor(80, 80, 80)

    def add_p(text, bold_prefix=None, space_after=6, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = FONT_FAMILY
            rb.font.size = Pt(12)
            rb.font.bold = True
            rb.font.color.rgb = COLOR_TEXT
        r = p.add_run(text)
        r.font.name = FONT_FAMILY
        r.font.size = Pt(12)
        r.font.italic = italic
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = FONT_FAMILY
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = FONT_FAMILY
        r.font.size = Pt(13.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = FONT_FAMILY
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_callout(quote_text, title=None):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F2F5F8")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=140)
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.line_spacing = 1.15
        if title:
            ra = cp.add_run(title + "\n")
            ra.font.name = FONT_FAMILY
            ra.font.size = Pt(11)
            ra.font.bold = True
            ra.font.color.rgb = COLOR_PRIMARY
        rq = cp.add_run(quote_text)
        rq.font.name = FONT_FAMILY
        rq.font.size = Pt(11)
        rq.font.italic = True
        rq.font.color.rgb = COLOR_TEXT
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_figure(img_path, caption, width_in=5.8):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run()
            run.add_picture(img_path, width=Inches(width_in))
            
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(12)
            c_run = cp.add_run(caption)
            c_run.font.name = FONT_FAMILY
            c_run.font.size = Pt(10)
            c_run.font.italic = True
            c_run.font.color.rgb = COLOR_MUTED

    # ==========================================
    # CARÁTULA OFICIAL
    # ==========================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(20)
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("UNIVERSIDAD NACIONAL DE SAN AGUSTÍN DE AREQUIPA")
    r_inst.font.name = FONT_FAMILY
    r_inst.font.size = Pt(16)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    p_fac = doc.add_paragraph()
    p_fac.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fac.paragraph_format.space_after = Pt(20)
    r_fac = p_fac.add_run("FACULTAD DE INGENIERÍA DE PRODUCCIÓN Y SERVICIOS\nESCUELA PROFESIONAL DE INGENIERÍA DE SISTEMAS")
    r_fac.font.name = FONT_FAMILY
    r_fac.font.size = Pt(13)
    r_fac.font.bold = True
    r_fac.font.color.rgb = COLOR_SECONDARY

    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_before = Pt(30)
    p_tit.paragraph_format.space_after = Pt(8)
    r_tit = p_tit.add_run("INFORME TÉCNICO DE INGENIERÍA DE SOFTWARE")
    r_tit.font.name = FONT_FAMILY
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(40)
    r_sub = p_sub.add_run("ARQUITECTURA DE SOFTWARE C4 Y PATRONES DE DISEÑO GOF\nEN LA PLATAFORMA INTEGRAL «RUTACOLCA»\nCasos Prácticos AG-08 y AG-09 (Semanas 08 y 09)")
    r_sub.font.name = FONT_FAMILY
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Asignatura:", "Construcción de Software (Código: 1703243) — Semestre 2026-B"),
        ("Docente:", "Mg. Antonio Arroyo Paz (aarroyop@unsa.edu.pe)"),
        ("Competencias ABET:", "C.c (Diseño con restricciones realistas) · C.d (Trabajo en equipo) · C.g (Comunicación)"),
        ("Repositorio GitHub:", "https://github.com/leonardobaca7/rutaColca_C4.git"),
        ("Ciudad y Fecha:", "Arequipa, Octubre de 2026")
    ]
    for i, (k, v) in enumerate(meta_data):
        c0, c1 = meta_table.cell(i, 0), meta_table.cell(i, 1)
        c0.width, c1.width = Inches(2.2), Inches(4.3)
        set_cell_background(c0, "F5F7FA")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)
        
        pk = c0.paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.font.name = FONT_FAMILY
        rk.font.size = Pt(11)
        rk.font.bold = True
        rk.font.color.rgb = COLOR_PRIMARY
        
        pv = c1.paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.name = FONT_FAMILY
        rv.font.size = Pt(11)
        rv.font.color.rgb = COLOR_TEXT

    p_team_tit = doc.add_paragraph()
    p_team_tit.paragraph_format.space_before = Pt(30)
    p_team_tit.paragraph_format.space_after = Pt(6)
    r_team_tit = p_team_tit.add_run("Autores e Integrantes del Equipo:")
    r_team_tit.font.name = FONT_FAMILY
    r_team_tit.font.size = Pt(12)
    r_team_tit.font.bold = True
    r_team_tit.font.color.rgb = COLOR_PRIMARY

    team_table = doc.add_table(rows=5, cols=4)
    team_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Estudiante", "Correo Institucional", "Rol AG-08 (C4)", "Rol AG-09 (Patrones)"]
    for j, h in enumerate(headers):
        cell = team_table.cell(0, j)
        set_cell_background(cell, "0B3D62")
        set_cell_margins(cell, 80, 80, 90, 90)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.name = FONT_FAMILY
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    team_data = [
        ("Emanuel Hilacondo Begazo", "ehilacondob@unsa.edu.pe", "Modelador / Revisor", "Analista de Diseño"),
        ("Leonardo Juan José Baca Calsin", "lbacac@unsa.edu.pe", "Redactor DSL / Vocero", "Líder Técnico e Implementación"),
        ("Mauro Snayder Sullca Mamani", "msullcam@unsa.edu.pe", "Revisor / Redactor DSL", "Evaluador de Arquitectura"),
        ("German Arturo Chipana Jeronimo", "gchipanaj@unsa.edu.pe", "Vocero / Modelador", "Diseñador de Modelos UML")
    ]
    for row_idx, data in enumerate(team_data, start=1):
        bg = "FFFFFF" if row_idx % 2 == 1 else "F9FAFC"
        for col_idx, val in enumerate(data):
            cell = team_table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = FONT_FAMILY
            r.font.size = Pt(10)
            if col_idx == 0:
                r.font.bold = True
            r.font.color.rgb = COLOR_TEXT

    doc.add_page_break()

    # ==========================================
    # RESUMEN EJECUTIVO
    # ==========================================
    add_heading_1("Resumen Ejecutivo")
    add_p(
        "El presente informe consolida el trabajo arquitectónico y de diseño orientado a objetos desarrollado durante "
        "las Semanas 08 y 09 para la plataforma «RutaColca», sistema concebido para optimizar la coordinación y venta de "
        "tours turísticos hacia el Cañón del Colca (Arequipa y Chivay). En una primera fase (AG-08), se aplicó el modelo C4 "
        "de Simon Brown junto con la filosofía de Diagramas como Código (DaC) sobre Structurizr DSL, modelando cuatro vistas "
        "canónicas: Contexto, Contenedores Lógicos, Dinámico y Topología de Despliegue en Producción, complementado con el registro "
        "del ADR-002 para trabajo sin conexión mediante IndexedDB y mitigación de sobreventa."
    )
    add_p(
        "En la segunda fase (AG-09), el diseño se profundizó hacia el nivel de componentes de software orientados a objetos, abordando "
        "cinco escenarios de crecimiento comercial. Se aplicaron cuatro patrones de diseño GoF: Factory Method y Adapter (para cobros "
        "multipasarela con Tarjeta, Yape y PayPal), Strategy (para tarifas dinámicas de temporada, estudiantes y grupos), Observer "
        "(para notificaciones concurrentes a turistas, guías y agencias) y Adapter (para integrar la API legada de AUTOCOLCA). "
        "Crucialmente, el equipo resolvió el reto de la «paternite» en el quinto escenario demostrando con criterio de ingeniería "
        "por qué NO debe forzarse un patrón Proxy artesanal cuando un índice B-Tree relacional en PostgreSQL 16 y cabeceras de caché HTTP "
        "resuelven el problema con mayor confiabilidad y menor costo de mantenimiento."
    )

    # ==========================================
    # SECCIÓN 1: INTRODUCCIÓN
    # ==========================================
    add_heading_1("1. Introducción y Contexto del Caso RutaColca")
    add_p(
        "El Valle del Colca recibe anualmente a miles de visitantes nacionales y extranjeros atraídos por atractivos como la "
        "Cruz del Cóndor, los baños termales y rutas de trekking. A pesar de este dinamismo, las agencias de turismo locales "
        "en Arequipa y Chivay gestionan sus operaciones mediante canales informales: llamadas telefónicas, mensajes no coordinados "
        "en WhatsApp y anotaciones físicas en cuadernos de ventanilla. Esta operativa manual genera descoordinación logística con los "
        "guías, sobreventa de cupos en temporada alta y dificultades para validar los boletos turísticos obligatorios de AUTOCOLCA."
    )
    add_p(
        "Para transformar este ecosistema, un consorcio de agencias solicitó el diseño integral de RutaColca. El sistema debe responder "
        "a seis historias de usuario clave en su MVP (HU-01 a HU-06), respetando las restricciones de tiempo y tecnología del ADR-001 "
        "(React SPA, React PWA, Spring Boot en Java 21, PostgreSQL 16 y Docker en la nube, entregado en 1 mes por un equipo de 4 personas)."
    )

    # ==========================================
    # SECCIÓN 2: INVENTARIO C4
    # ==========================================
    add_heading_1("2. Inventario Arquitectónico y Clasificación C4")
    add_p(
        "De acuerdo con el modelo C4 de Simon Brown, el diseño comenzó por la delimitación estricta de responsabilidades entre "
        "personas, sistemas propios, sistemas externos y contenedores lógicos, evitando mezclar elementos de infraestructura técnica."
    )

    inv_table = doc.add_table(rows=14, cols=3)
    inv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    inv_headers = ["Elemento del Caso", "Tipo C4 Formal", "Justificación Arquitectónica"]
    for j, h in enumerate(inv_headers):
        cell = inv_table.cell(0, j)
        set_cell_background(cell, "0B3D62")
        set_cell_margins(cell, 70, 70, 80, 80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.name = FONT_FAMILY
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    inv_data = [
        ("Turista", "Persona", "Usuario final que busca tours, selecciona fechas, reserva y paga en línea."),
        ("Operador de agencia", "Persona", "Personal de ventanilla en Arequipa o Chivay que publica itinerarios y gestiona cupos."),
        ("Guía de turismo", "Persona", "Profesional de campo que recibe listas e itinerarios por mensajería sin entrar a la web."),
        ("Administrador", "Persona", "Auditor interno que autoriza nuevas agencias y monitorea reportes financieros."),
        ("RutaColca", "Sistema de Software (Propio)", "Núcleo de software en estudio que centraliza catálogo, reservas, pagos y validación."),
        ("Pasarela de pagos", "Sistema de Software (Externo)", "Proveedor financiero (Niubiz/Stripe) que procesa pagos con tarjeta y billeteras móviles."),
        ("Sistema de AUTOCOLCA", "Sistema de Software (Externo)", "Servicio gubernamental externo consultado vía HTTPS para autenticar boletos turísticos."),
        ("WhatsApp Business", "Sistema de Software (Externo)", "Plataforma de mensajería empresarial para entrega automatizada de manifiestos."),
        ("Web del turista", "Contenedor (Aplicación)", "Single Page Application (React) para consulta de catálogo, reserva y pago digital."),
        ("Portal de agencias", "Contenedor (Aplicación)", "Progressive Web App (React PWA) para gestión operativa y ventas con o sin conexión."),
        ("API Backend", "Contenedor (Aplicación)", "Servicio REST central (Java 21, Spring Boot) que orquesta lógica y datos."),
        ("Base de Datos Principal", "Contenedor (Almacén de Datos)", "Motor relacional PostgreSQL 16 para persistencia ACID de transacciones."),
        ("Almacén Local de Reservas", "Contenedor (Almacén de Datos)", "Base de datos IndexedDB en el navegador para registrar ventas offline en Chivay.")
    ]
    for row_idx, (e_name, e_type, e_just) in enumerate(inv_data, start=1):
        bg = "FFFFFF" if row_idx % 2 == 1 else "F9FAFC"
        for col_idx, text_val in enumerate([e_name, e_type, e_just]):
            cell = inv_table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 60, 60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text_val)
            r.font.name = FONT_FAMILY
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.font.bold = True
            elif col_idx == 1:
                r.font.color.rgb = COLOR_SECONDARY
                r.font.bold = True
            else:
                r.font.color.rgb = COLOR_TEXT

    add_callout(
        "Decisión sobre el Guía de Turismo: Aunque el guía no tiene credenciales en el portal web, es una Persona en el "
        "modelo C4 porque es el receptor humano final del manifiesto de viaje mediante WhatsApp. Omitirlo rompería la trazabilidad del servicio.\n"
        "Exclusión de Docker: Docker es un motor de virtualización de infraestructura de despliegue, no un contenedor conceptual en C4. "
        "Confundir ambos es un error penalizado en la rúbrica.",
        title="Decisiones Críticas del Inventario:"
    )

    # ==========================================
    # SECCIÓN 3: MODELADO C4 (AG-08)
    # ==========================================
    add_heading_1("3. Modelado Arquitectónico C4 de RutaColca (AG-08)")
    add_p(
        "A partir de `workspace.dsl`, el compilador Structurizr CLI proyectó las vistas canónicas del sistema, garantizando "
        "consistencia absoluta entre niveles de zoom y relaciones unidireccionales explícitas."
    )

    add_heading_2("3.1. Diagrama de Contexto del Sistema (Nivel 1)")
    add_p(
        "Muestra a RutaColca como caja única delimitada por las cuatro personas interesadas y los tres sistemas externos. "
        "Todas las relaciones emplean verbos activos y específicos con su respectivo protocolo."
    )
    add_figure(
        "docs/architecture/c4/ruta-colca/export/structurizr-Contexto.png",
        "Figura 1. Diagrama de Contexto de la plataforma RutaColca (Structurizr DSL y PlantUML).",
        width_in=5.8
    )

    add_heading_2("3.2. Diagrama de Contenedores (Nivel 2 · Operación Estándar)")
    add_p(
        "Detalla las cuatro unidades de ejecución según el ADR-001: Web del Turista (React SPA), Portal de Agencias (React PWA), "
        "API Backend (Java 21, Spring Boot) y Base de Datos (PostgreSQL 16), con protocolos JSON/HTTPS y JDBC."
    )
    add_figure(
        "docs/architecture/c4/ruta-colca/export/structurizr-Contenedores.png",
        "Figura 2. Diagrama de Contenedores de RutaColca en operación estándar conectada.",
        width_in=5.2
    )

    add_heading_2("3.3. Diagrama Dinámico: Reservar y Pagar un Tour (HU-02 y HU-03)")
    add_p(
        "Describe el flujo de 8 pasos para procesar una reserva en línea. El equipo estableció y justificó la decisión de negocio:",
        bold_prefix="Orquestación Temporal: "
    )
    add_callout(
        "¿Se debe validar el boleto con AUTOCOLCA antes o después de cobrar en la pasarela de pagos?\n"
        "Decisión: Se valida OBLIGATORIAMENTE ANTES de cobrar. Si se debitaran los fondos y luego AUTOCOLCA rechazara el boleto por "
        "encontrarse vencido o adulterado, la empresa asume costos financieros por contracargos bancarios no reembolsables y fricción operativa "
        "al gestionar devoluciones manuales. Validar primero protege financieramente a la plataforma y garantiza la validez legal del tour.",
        title="Decisión de Diseño Dinámico:"
    )
    add_figure(
        "docs/architecture/c4/ruta-colca/export/structurizr-Dinamico.png",
        "Figura 3. Diagrama Dinámico: Interacción secuencial numerada para reserva y pago en línea.",
        width_in=6.0
    )

    add_heading_2("3.4. Diagrama de Despliegue en Producción")
    add_p(
        "Especifica la topología de cómputo en la nube para soportar los picos estacionales de julio y agosto: Balanceador de Carga "
        "con TLS (`AWS ALB / Nginx`), dos máquinas virtuales Ubuntu 24.04 con Docker para escalabilidad de la API y servicio gestionado "
        "de base de datos PostgreSQL 16."
    )
    add_figure(
        "docs/architecture/c4/ruta-colca/export/structurizr-DespliegueProduccion.png",
        "Figura 4. Diagrama de Despliegue en Producción con clúster balanceado y alta disponibilidad.",
        width_in=5.5
    )

    # ==========================================
    # SECCIÓN 4: RETO OFFLINE Y ADR-002
    # ==========================================
    add_heading_1("4. Resolución del Reto Sin Conexión en Chivay (HU-05) y ADR-002")
    add_p(
        "Para responder a las caídas de internet en Chivay sin violar la restricción de 1 mes de desarrollo, se diseñó una solución "
        "Offline-First basada en IndexedDB en el navegador del operador y la política de Partición de Cupos Garantizados documentada en el ADR-002."
    )
    add_figure(
        "docs/architecture/c4/ruta-colca/export/structurizr-ContenedoresReto.png",
        "Figura 5. Diagrama de Contenedores con soporte offline en Chivay (`ContenedoresReto`).",
        width_in=5.5
    )

    # ==========================================
    # SECCIÓN 5: PATRONES DE DISEÑO GOF (AG-09)
    # ==========================================
    add_heading_1("5. Diseño Orientado a Objetos: Patrones de Diseño GoF (AG-09)")
    add_p(
        "En la Semana 09, RutaColca profundizó su diseño hacia el nivel de componentes internos y clases en Java 21, resolviendo "
        "cinco escenarios de crecimiento mediante el catálogo GoF y principios de diseño SOLID."
    )

    add_heading_2("5.1. Escenario 1: Procesamiento Multipasarela (Factory Method + Adapter)")
    add_p(
        "Para aceptar simultáneamente Tarjeta de Crédito, Yape y PayPal sin que `ReservaService` dependa de clases concretas, se implementó "
        "el patrón Factory Method combinado con Adaptadores. La fábrica `PasarelaPagoFactory` decide qué adaptador instanciar, mientras "
        "que el cliente solo conoce el contrato polimórfico `PasarelaPago` (Principio Abierto/Cerrado y Programar para Interfaces)."
    )
    add_figure(
        "docs/patterns/ruta-colca/diagramas/patron-pagos-factory-adapter.png",
        "Figura 6. Diagrama de Clases: Patrón Factory Method y Adaptadores para pasarelas de pago.",
        width_in=5.6
    )

    add_heading_2("5.2. Escenario 2: Tarifas Dinámicas (Strategy)")
    add_p(
        "El costo del tour varía según temporada alta (+25%), descuento grupal (-15% a partir de 5 pasajeros) o condición estudiantil (-30%). "
        "El patrón Strategy encapsula cada algoritmo de cálculo en clases concretas (`PrecioTemporadaAltaStrategy`, `PrecioGrupoStrategy`, "
        "`PrecioEstudianteStrategy`). La clase `Reserva` actúa como contexto y delega el cálculo, permitiendo intercambiar estrategias en caliente."
    )
    add_figure(
        "docs/patterns/ruta-colca/diagramas/patron-tarifas-strategy.png",
        "Figura 7. Diagrama de Clases: Patrón Strategy para tarificación dinámica de tours.",
        width_in=5.6
    )

    add_heading_2("5.3. Escenario 3: Notificaciones Concurrentes Desacopladas (Observer)")
    add_p(
        "Al confirmar una reserva se debe alertar al turista (correo), al guía (WhatsApp) y a la agencia (tablero PWA). El patrón Observer "
        "desacopla el servicio de reservas de los canales de comunicación: `PublicadorReservas` emite el evento inmutable `EventoReservaConfirmada` "
        "hacia los observadores suscritos de forma independiente (Bajo acoplamiento y Responsabilidad Única)."
    )
    add_figure(
        "docs/patterns/ruta-colca/diagramas/patron-notificaciones-observer.png",
        "Figura 8. Diagrama de Clases: Patrón Observer para despacho de notificaciones multicanal.",
        width_in=5.6
    )

    add_heading_2("5.4. Escenario 4: Integración con la API Legada de AUTOCOLCA (Adapter)")
    add_p(
        "La entidad estatal AUTOCOLCA expone un servicio con firmas y estructuras propietarias (`consultarVigenciaBoletoServicioPublico`). "
        "El patrón Adapter implementa la interfaz de dominio `ValidadorBoleto` a través de `AutocolcaApiAdapter`, protegiendo el núcleo de la "
        "aplicación de cambios en el proveedor externo (Capa Anticorrupción / ACL e Inversión de Dependencias)."
    )
    add_figure(
        "docs/patterns/ruta-colca/diagramas/patron-autocolca-adapter.png",
        "Figura 9. Diagrama de Clases: Patrón Adapter como Capa Anticorrupción frente a AUTOCOLCA.",
        width_in=5.4
    )

    add_heading_2("5.5. Escenario 5: El Reto de la «Paternite» (Consultas Masivas de Disponibilidad)")
    add_p(
        "La disponibilidad de cupos se consulta cientos de veces por minuto. Ante la tentación de forzar un patrón estructural Proxy de caché "
        "en memoria, el equipo aplicó el principio de Ousterhout:",
        bold_prefix="Defensa contra la Paternite: "
    )
    add_callout(
        "¿Por qué NO usar un patrón Proxy en el Escenario 5?\n"
        "1. Inconsistencia y Overbooking: Un Proxy en memoria en un clúster con 2 servidores cloud genera copias locales desincronizadas, "
        "permitiendo que dos agencias vendan el último asiento.\n"
        "2. Complejidad innecesaria: Requiere código multihilo complejo y gestión manual de invalidación.\n"
        "3. Solución nativa superior: Un índice relacional B-Tree en PostgreSQL 16 sobre (tour_id, fecha_salida) resuelve las lecturas en menos de "
        "1 milisegundo directamente desde el motor ACID. Adicionalmente, cabeceras HTTP Cache-Control: max-age=30 en el balanceador absorben "
        "las peticiones repetidas sin añadir una sola línea de código espagueti en Java.",
        title="Dictamen Técnico sobre el Reto de la Paternite:"
    )

    # ==========================================
    # SECCIÓN 6: AUDITORÍA Y TRAZABILIDAD IA
    # ==========================================
    add_heading_1("6. Gobernanza de Inteligencia Artificial y Auditoría Cruzada")
    add_p(
        "Bajo la directiva de la asignatura («La IA propone, el equipo decide y verifica»), el equipo documentó todas las interacciones "
        "con asistentes inteligentes, identificando y rechazando sistemáticamente las alucinaciones y la tendencia de los LLM al sobrediseño."
    )

    ia_table = doc.add_table(rows=6, cols=4)
    ia_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ia_headers = ["Actividad", "Propuesta de la IA", "Evaluación Crítica del Equipo", "Decisión Final"]
    for j, h in enumerate(ia_headers):
        cell = ia_table.cell(0, j)
        set_cell_background(cell, "0B3D62")
        set_cell_margins(cell, 70, 70, 80, 80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.name = FONT_FAMILY
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    ia_records = [
        ("AG-08 · Paso 1", "Incluir Docker como contenedor lógico C4.", "Violación conceptual grave de la regla de oro C4.", "Rechazado. Reubicado en el diagrama de despliegue."),
        ("AG-08 · Paso 3", "Incorporar Kafka y Redis al backend.", "Sobrediseño que excede el plazo del MVP de 1 mes.", "Rechazado. Se mantuvo el stack del ADR-001."),
        ("AG-08 · Paso 4", "Cobrar antes de validar con AUTOCOLCA.", "Riesgo financiero severo de contracargos bancarios.", "Reordenado. Validación previa obligatoria."),
        ("AG-09 · Paso 1", "Usar Abstract Factory en pasarelas de pago.", "No existen múltiples familias abstractas.", "Ajustado a Factory Method + Adapter."),
        ("AG-09 · Reto", "Crear Proxy de caché en memoria para cupos.", "Caso clásico de paternite; riesgo de sobreventa.", "Rechazado. Resuelto con índice en BD y HTTP Cache.")
    ]
    for row_idx, data_row in enumerate(ia_records, start=1):
        bg = "FFFFFF" if row_idx % 2 == 1 else "F9FAFC"
        for col_idx, text_val in enumerate(data_row):
            cell = ia_table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 60, 60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text_val)
            r.font.name = FONT_FAMILY
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.font.bold = True
            elif col_idx == 3:
                r.font.color.rgb = COLOR_SECONDARY
                r.font.bold = True
            r.font.color.rgb = COLOR_TEXT

    # ==========================================
    # SECCIÓN 7: RESPUESTAS DE EVALUACIÓN DOCENTE
    # ==========================================
    add_heading_1("7. Respuestas Argumentadas para la Evaluación Docente")
    add_p(
        "1. ¿Por qué RutaColca aparece como una sola caja en el contexto y como un límite con contenedores en el siguiente nivel?\n"
        "Porque atienden a propósitos y audiencias diferentes. El contexto oculta los detalles técnicos para dialogar con gerentes de agencias "
        "y autoridades de turismo. El diagrama de contenedores hace un zoom in para el equipo técnico, mostrando las piezas ejecutables (SPA, PWA, "
        "API Spring Boot, PostgreSQL) y los protocolos de integración.\n\n"
        "2. ¿Qué diagrama le mostrarían a la gerente de una agencia y cuál al equipo de operaciones?\n"
        "A la gerente le mostraríamos el Diagrama de Contexto (Nivel 1), que explica qué servicios ofrece la plataforma y cómo se conecta sin jerga "
        "técnica. Al equipo de operaciones le mostraríamos el Diagrama de Despliegue, que detalla las máquinas virtuales, balanceadores y base de datos gestionada.\n\n"
        "3. ¿Qué propuesta de la IA rechazaron y qué regla de C4 usaron para hacerlo?\n"
        "Rechazamos la propuesta de considerar a Docker como contenedor C4. Aplicamos la regla oficial de Simon Brown: 'un contenedor C4 es una aplicación "
        "o almacén de datos propio, no una tecnología de virtualización de infraestructura'.\n\n"
        "4. ¿Qué cambió en su modelo al resolver el reto sin conexión y en qué diagramas se nota?\n"
        "Se agregó el contenedor de datos local IndexedDB dentro del Portal PWA y relaciones de sincronización por lotes. Se refleja en la vista "
        "ContenedoresReto y en el Diagrama de Despliegue dentro del nodo Tablet de Agencia.\n\n"
        "5. Ticket de salida de Patrones: Un banco nuevo ofrece pagos con su propia API. ¿Qué patrón usaría y qué clase crearía?\n"
        "Usaríamos el patrón Adaptador (Adapter) creando una nueva clase concreta, por ejemplo BancoNuevoAdapter, que implemente la interfaz "
        "PasarelaPago y traduzca las llamadas hacia el SDK del banco, respetando el principio Abierto/Cerrado (OCP)."
    )

    # ==========================================
    # SECCIÓN 8: CONCLUSIONES Y REFERENCIAS
    # ==========================================
    add_heading_1("8. Conclusiones y Referencias Bibliográficas")
    add_p("1. Continuidad entre arquitectura y diseño: El modelo C4 proporcionó la visión macroscópica del sistema, mientras que los patrones GoF estructuraron las clases internas garantizando bajo acoplamiento y alta cohesión.", bold_prefix="• ")
    add_p("2. Sentido crítico frente a la 'paternite': El verdadero mérito del diseño no radica en sobrecargar el código con patrones, sino en discernir cuándo la infraestructura relacional nativa es más simple y robusta.", bold_prefix="• ")
    add_p("3. Trazabilidad con Diagramas como Código: Structurizr DSL y PlantUML aseguran que la documentación viva evolucione al ritmo del código sin desincronizaciones.", bold_prefix="• ")

    add_heading_2("Referencias Bibliográficas (Norma APA 7ma Edición)")
    add_p("Bass, L., Clements, P., & Kazman, R. (2021). Software architecture in practice (4th ed.). Addison-Wesley.")
    add_p("Brown, S. (2026). The C4 model for visualising software architecture. O’Reilly Media / c4model.com.")
    add_p("Gamma, E., Helm, R., Johnson, R., & Vlissides, J. (1994). Design patterns: Elements of reusable object-oriented software. Addison-Wesley.")
    add_p("ISO/IEC/IEEE. (2022). ISO/IEC/IEEE 42010:2022 Software, systems and enterprise — Architecture description. International Organization for Standardization.")
    add_p("Ousterhout, J. (2018). A philosophy of software design. Yaknyam Press.")
    add_p("Valente, M. T. (2020). Engenharia de software moderna (cap. 5: Princípios de projeto; cap. 6: Padrões de projeto). Independente. https://engsoftmoderna.info/")

    output_docx = "docs/INFORME_TECNICO_RUTACOLCA_AG08_AG09.docx"
    doc.save(output_docx)
    print(f"Report successfully saved: {output_docx}")

if __name__ == "__main__":
    generate_report()
